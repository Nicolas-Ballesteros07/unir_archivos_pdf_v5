import os
import zipfile
import tempfile
import shutil
from typing import List, Dict
from pypdf import PdfWriter, PdfReader
from openpyxl import load_workbook
from PIL import Image
import re
import xlrd

IMAGE_EXTENSIONS = {'.jpg', '.jpeg', '.png', '.gif', '.bmp', '.tiff', '.tif', '.webp'}

def merge_pdfs(pdf_paths: List[str], output_path: str) -> None:
    if not pdf_paths:
        raise ValueError("No se proporcionaron archivos PDF")
    writer = PdfWriter()
    for path in pdf_paths:
        reader = PdfReader(path)
        for page in reader.pages:
            writer.add_page(page)
    with open(output_path, "wb") as f_out:
        writer.write(f_out)

def process_zip_to_pdf(zip_path: str, output_path: str) -> None:
    temp_dir = tempfile.mkdtemp()
    try:
        with zipfile.ZipFile(zip_path, 'r') as zf:
            zf.extractall(temp_dir)

        pdf_files = []
        image_files = []

        for root, dirs, files in os.walk(temp_dir):
            for file in files:
                full_path = os.path.join(root, file)
                ext = os.path.splitext(file)[1].lower()
                if ext == '.pdf':
                    pdf_files.append(full_path)
                elif ext in IMAGE_EXTENSIONS:
                    try:
                        with Image.open(full_path) as img:
                            img.verify()
                        image_files.append(full_path)
                    except Exception:
                        continue

        if not pdf_files and not image_files:
            raise ValueError("El ZIP no contiene archivos PDF ni imágenes válidas.")

        pdf_files.sort()
        image_files.sort()

        temp_pdf_paths = []
        for img_path in image_files:
            try:
                with Image.open(img_path) as img:
                    if img.mode in ('RGBA', 'LA', 'P'):
                        img = img.convert('RGB')
                    temp_img_pdf = os.path.join(temp_dir, f"__img_{len(temp_pdf_paths)}.pdf")
                    img.save(temp_img_pdf, "PDF")
                    temp_pdf_paths.append(temp_img_pdf)
            except Exception:
                continue

        all_pdfs = pdf_files + temp_pdf_paths
        merge_pdfs(all_pdfs, output_path)

    finally:
        shutil.rmtree(temp_dir, ignore_errors=True)

def read_excel_payments(file_path: str) -> List[Dict[str, str]]:
    ext = os.path.splitext(file_path)[1].lower()

    if ext == '.xlsx':
        wb = load_workbook(file_path, data_only=True)
        if 'Pagos Totales' not in wb.sheetnames:
            raise ValueError("El archivo Excel no contiene la hoja 'Pagos Totales'")
        ws = wb['Pagos Totales']

        if str(ws['C2'].value).strip() != 'Proveedor':
            raise ValueError("La celda C2 no contiene 'Proveedor'")
        if str(ws['D2'].value).strip() != 'Factura':
            raise ValueError("La celda D2 no contiene 'Factura'")

        data = []
        for row in ws.iter_rows(min_row=3, max_row=ws.max_row, min_col=3, max_col=4):
            proveedor = row[0].value
            factura = row[1].value
            if proveedor is not None and factura is not None:
                proveedor_str = str(proveedor).strip()
                # Si factura es numérico, convertir a entero si es un entero exacto
                if isinstance(factura, float) and factura == int(factura):
                    factura_str = str(int(factura))
                else:
                    factura_str = str(factura).strip()
                data.append({
                    'proveedor': proveedor_str,
                    'factura': factura_str
                })
        return data

    elif ext == '.xls':
        wb = xlrd.open_workbook(file_path)
        if 'Pagos Totales' not in wb.sheet_names():
            raise ValueError("El archivo Excel no contiene la hoja 'Pagos Totales'")
        ws = wb.sheet_by_name('Pagos Totales')

        if ws.cell_value(1, 2) != 'Proveedor':
            raise ValueError("La celda C2 no contiene 'Proveedor'")
        if ws.cell_value(1, 3) != 'Factura':
            raise ValueError("La celda D2 no contiene 'Factura'")

        data = []
        for row_idx in range(2, ws.nrows):
            proveedor = ws.cell_value(row_idx, 2)
            factura = ws.cell_value(row_idx, 3)
            if proveedor and factura:
                proveedor_str = str(proveedor).strip()
                # Limpiar .0 en números enteros
                if isinstance(factura, float) and factura == int(factura):
                    factura_str = str(int(factura))
                else:
                    factura_str = str(factura).strip()
                data.append({
                    'proveedor': proveedor_str,
                    'factura': factura_str
                })
        return data

    else:
        raise ValueError("Formato de archivo no soportado. Use .xls o .xlsx")

def sanitize_filename(name: str) -> str:
    return re.sub(r'[\\/*?:"<>|]', "", name).strip()